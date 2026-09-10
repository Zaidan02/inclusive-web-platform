from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(sys.argv[1] if len(sys.argv) > 1 else "JoIn_Final_Test_Register.xlsx")

NAVY = "0B234A"
BLUE = "1769D2"
GREEN = "16835A"
LIGHT_GREEN = "E8F6EE"
RED = "B42318"
LIGHT_RED = "FDECEC"
ORANGE = "B24B2A"
LIGHT_ORANGE = "FFF4E8"
PURPLE = "5B3A4E"
LIGHT_BLUE = "EAF2FD"
GREY = "E8EDF4"
TEXT = "17233D"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="CBD5E1")


AREAS = [
    ("AI Profile Suggestions", 60, "Development Team", ">=95% precision and recall; 0 unsafe inference"),
    ("Voice Transcription", 30, "Development Team", ">=90% success; >=95% critical-value accuracy"),
    ("Voice Navigation", 60, "Development Team", ">=90% action success; 0 unauthorized actions"),
    ("Matching Engine", 45, "Development Team", "100% formula and eligibility accuracy"),
    ("Complete Workflows", 20, "Development Team", "100% critical flows; >=95% overall"),
    ("Security and Privacy", 20, "Development Team", "100%; no critical or high finding open"),
    ("Accessibility", 18, "Development Team", "No critical blockers; 100% keyboard completion"),
    ("Multilingual and RTL", 18, "Development Team", "100%; no missing keys or blocking RTL defects"),
    ("Browser and Device", 12, "Development Team", ">=90%; 100% critical flows"),
    ("Performance and Reliability", 27, "Development Team", "<1% 5xx; agreed p95 thresholds met"),
]

LANGUAGES = {"en": "English", "fr": "French", "ar": "Arabic"}


def case(test_id, area, subsystem, owner, priority, language, condition, preconditions, scenario, steps, expected):
    return [
        test_id, area, subsystem, owner, priority, language, condition, preconditions,
        scenario, steps, expected, "", "", "", "", "", "", "", "", "", "",
    ]


def ai_cases():
    personas = [
        ("Complete profile", "All profile fields, high-school education, three independent abilities, both opportunity types, Cashier interest, explicit Hand selection and at least three task skills.", "Every explicitly stated field is suggested; no value is saved before confirmation."),
        ("Basic information only", "First name, last name, phone, location and a candidate-provided About sentence only.", "Only supported explicitly stated profile fields are suggested."),
        ("Independent foundation abilities", "Candidate explicitly says they read, write and count independently.", "All three labelled abilities map to independent."),
        ("Abilities with support", "Candidate explicitly needs support for reading, writing and counting.", "All three labelled abilities map to with_support without exclusion."),
        ("Mixed and not-yet abilities", "Reading independently, writing with support and counting not yet.", "Each ability maps to its own correct controlled value."),
        ("Work only", "Candidate explicitly wants paid work and names a controlled position.", "Opportunity preference maps to work and the controlled position is suggested."),
        ("Training only", "Candidate explicitly wants a hospitality training program.", "Opportunity preference maps to training without inventing job experience."),
        ("Work and training", "Candidate explicitly wants both work and hospitality training.", "Opportunity preference maps to both."),
        ("Explicit disability and skills", "Candidate explicitly names a catalogue disability and gives evidence for controlled tasks.", "Only the explicitly named disability and evidence-backed catalogue tasks are suggested."),
        ("Ambiguous safety boundary", "Narrative is vague about health, abilities and job experience.", "No disability diagnosis, unsupported ability or invented catalogue task is suggested; ambiguity remains unmapped."),
    ]
    rows = []
    for code, language in LANGUAGES.items():
        for quality_code, quality in (("C", "Clean text"), ("I", "Imperfect transcript")):
            for index, (name, scenario, expected) in enumerate(personas, 1):
                rows.append(case(
                    f"AI-{code.upper()}-{quality_code}{index:02d}", "AI Profile Suggestions", name,
                    "Development Team", "Critical" if index in (1, 9, 10) else "High", language, quality,
                    "Candidate fixture authenticated; local or staging AI service healthy; profile state recorded.",
                    scenario,
                    "Reset/record profile state; enter the prepared narrative; select the declared input language; provide consent; create suggestions; compare every item with the gold expectation; confirm selected valid items where required.",
                    expected,
                ))
    return rows


def transcription_cases():
    contents = [
        ("Identity and contact", "Speak a fictional name, Beirut location and an eight-digit phone number.", "Transcript preserves the name, location and all digits correctly."),
        ("Education", "State a controlled education level and one work goal.", "Education and goal words are transcribed without changing their meaning."),
        ("Foundation abilities", "State reading, writing and counting ability in one natural sentence.", "All three ability concepts appear correctly in the editable transcript."),
        ("Opportunity preference", "State work, training or both and name Cashier as the preferred position.", "Opportunity type and Cashier interest are preserved."),
        ("Disability and task evidence", "Explicitly name Hand and three Cashier task skills.", "The explicit term and important task phrases remain understandable and editable."),
    ]
    rows = []
    for code, language in LANGUAGES.items():
        for condition_code, condition in (("Q", "Quiet"), ("N", "Moderate noise / imperfect pronunciation")):
            for index, (name, scenario, expected) in enumerate(contents, 1):
                rows.append(case(
                    f"TR-{code.upper()}-{condition_code}{index:02d}", "Voice Transcription", name,
                    "Development Team", "High", language, condition,
                    "Microphone permission granted; correct input language selected; fictional test data only.",
                    scenario,
                    "Start recording; speak the prepared sentence once; stop recording; wait for transcription; compare expected and actual text; verify editing remains possible; record latency and evidence.",
                    expected,
                ))
    return rows


VOICE_COMMANDS = {
    "en": [
        ("Navigation", "Open jobs", "The Jobs view opens."),
        ("Navigation", "Open my profile", "The candidate Profile view opens."),
        ("Navigation", "Open my applications", "The Applications view opens."),
        ("Navigation", "Open privacy and data", "The Privacy and data view opens."),
        ("Navigation", "Go to the home page", "The public home page opens through a registered route."),
        ("Navigation", "Open voice navigation help", "The registered voice-help page opens."),
        ("Interface action", "Select work and training", "The Both opportunity option is selected and announced."),
        ("Interface action", "Set reading to independently", "Reading is set to independent and announced for review."),
        ("Interface action", "Choose Cashier", "The controlled Cashier option is selected or focused as designed."),
        ("Interface action", "Scroll down", "The page scrolls by the allowed bounded amount."),
        ("Interface action", "Clear the disability search", "Only the registered search field is cleared."),
        ("Interface action", "Save my profile", "The registered save action occurs only when validation passes."),
        ("Page question", "What is this page for?", "A grounded answer describes only the current page."),
        ("Page question", "Which fields are required?", "A grounded answer identifies the current required fields."),
        ("Page question", "How is my match score calculated?", "A grounded explanation is returned without inventing a score."),
        ("Page question", "Is assistance available for this job?", "The answer uses only visible current-job information."),
        ("Confirmation", "Submit my application", "The action waits for explicit confirmation before submission."),
        ("Confirmation", "Reset my profile", "The action waits for explicit confirmation before reset."),
        ("Rejected", "Delete another candidate", "The unauthorized command is rejected and no state changes."),
        ("Rejected", "Open the secret management page", "The unregistered destination is rejected safely."),
    ],
    "fr": [
        ("Navigation", "Ouvre les offres", "La vue des offres s’ouvre."),
        ("Navigation", "Ouvre mon profil", "La vue du profil candidat s’ouvre."),
        ("Navigation", "Ouvre mes candidatures", "La vue des candidatures s’ouvre."),
        ("Navigation", "Ouvre confidentialité et données", "La vue Confidentialité et données s’ouvre."),
        ("Navigation", "Va à la page d’accueil", "La page d’accueil publique s’ouvre par une route enregistrée."),
        ("Navigation", "Ouvre l’aide de navigation vocale", "La page d’aide vocale enregistrée s’ouvre."),
        ("Interface action", "Sélectionne emploi et formation", "L’option Les deux est sélectionnée et annoncée."),
        ("Interface action", "Mets la lecture sur autonome", "La lecture est définie sur autonome et annoncée."),
        ("Interface action", "Choisis Caissier", "Le poste contrôlé Cashier est sélectionné ou ciblé."),
        ("Interface action", "Fais défiler vers le bas", "La page défile de la quantité autorisée."),
        ("Interface action", "Efface la recherche de handicap", "Seul le champ de recherche enregistré est effacé."),
        ("Interface action", "Enregistre mon profil", "L’action enregistrée se produit uniquement si la validation réussit."),
        ("Page question", "À quoi sert cette page ?", "Une réponse fondée décrit uniquement la page actuelle."),
        ("Page question", "Quels champs sont obligatoires ?", "La réponse identifie les champs obligatoires actuels."),
        ("Page question", "Comment mon score est-il calculé ?", "Une explication fondée est fournie sans inventer de score."),
        ("Page question", "Une assistance est-elle disponible pour ce poste ?", "La réponse utilise uniquement les informations visibles."),
        ("Confirmation", "Envoie ma candidature", "Une confirmation explicite est exigée avant l’envoi."),
        ("Confirmation", "Réinitialise mon profil", "Une confirmation explicite est exigée avant la réinitialisation."),
        ("Rejected", "Supprime un autre candidat", "La commande non autorisée est rejetée sans modification."),
        ("Rejected", "Ouvre la page secrète de gestion", "La destination non enregistrée est rejetée."),
    ],
    "ar": [
        ("Navigation", "افتح الوظائف", "تفتح صفحة الوظائف."),
        ("Navigation", "افتح ملفي الشخصي", "تفتح صفحة ملف المرشح."),
        ("Navigation", "افتح طلباتي", "تفتح صفحة الطلبات."),
        ("Navigation", "افتح الخصوصية والبيانات", "تفتح صفحة الخصوصية والبيانات."),
        ("Navigation", "اذهب إلى الصفحة الرئيسية", "تفتح الصفحة الرئيسية عبر مسار مسجل."),
        ("Navigation", "افتح مساعدة التنقل الصوتي", "تفتح صفحة مساعدة الصوت المسجلة."),
        ("Interface action", "اختر العمل والتدريب", "يتم تحديد خيار العمل والتدريب والإعلان عنه."),
        ("Interface action", "ضع القراءة على نعم باستقلالية", "يتم ضبط القراءة على مستقل والإعلان عنها."),
        ("Interface action", "اختر كاشير", "يتم تحديد أو تركيز وظيفة Cashier المضبوطة."),
        ("Interface action", "انزل إلى أسفل", "يتم تمرير الصفحة بالمقدار المسموح."),
        ("Interface action", "امسح بحث الإعاقة", "يتم مسح حقل البحث المسجل فقط."),
        ("Interface action", "احفظ ملفي", "يتم الحفظ فقط إذا نجح التحقق."),
        ("Page question", "ما هدف هذه الصفحة؟", "تصف الإجابة الصفحة الحالية فقط."),
        ("Page question", "ما هي الحقول المطلوبة؟", "تحدد الإجابة الحقول المطلوبة الحالية."),
        ("Page question", "كيف يتم حساب نسبة المطابقة؟", "يتم شرح الحساب دون اختراع نتيجة."),
        ("Page question", "هل توجد مساعدة لهذه الوظيفة؟", "تعتمد الإجابة على معلومات الوظيفة الظاهرة فقط."),
        ("Confirmation", "أرسل طلبي", "ينتظر النظام تأكيداً صريحاً قبل الإرسال."),
        ("Confirmation", "أعد ضبط ملفي", "ينتظر النظام تأكيداً صريحاً قبل إعادة الضبط."),
        ("Rejected", "احذف مرشحاً آخر", "يتم رفض الأمر غير المصرح ولا تتغير البيانات."),
        ("Rejected", "افتح صفحة الإدارة السرية", "يتم رفض الوجهة غير المسجلة بأمان."),
    ],
}


def voice_navigation_cases():
    rows = []
    for code, commands in VOICE_COMMANDS.items():
        for index, (category, command, expected) in enumerate(commands, 1):
            rows.append(case(
                f"VN-{code.upper()}-{index:02d}", "Voice Navigation", category, "Development Team",
                "Critical" if category in ("Confirmation", "Rejected") else "High",
                LANGUAGES[code], "Quiet first; repeat failed cases with moderate noise",
                "Correct role authenticated where required; microphone enabled; voice language selected.",
                command,
                "Speak the command once; record transcript, classification, authorization result, final UI action, feedback and latency; capture evidence for failures and sensitive commands.",
                expected,
            ))
    return rows


def matching_cases():
    roles = [
        "Bakery / Pastry Worker", "Cashier", "Chocolate / Confectionery Worker",
        "Cold Kitchen - Garde manger", "Cold Preparation: Butcher / Poultry",
        "Fish preparation", "Food service point of sales management", "Head Barman / Barman",
        "Host / Hostess", "Ice Cream Maker", "Point of sales manager - Self service",
        "Pot - Dish - Glass Washing", "Rotating Chef - Self service, Banquet & other",
        "Sommelier", "Vegetable Preparation",
    ]
    archetypes = [
        ("Independent", "Candidate independently meets configured tasks and abilities.", "Displayed score equals the manual weighted oracle; eligibility passes."),
        ("With support", "Candidate can perform selected tasks with offered assistance.", "The 0.75 feasibility factor is applied only where support is offered; explanation matches."),
        ("Ineligible gate", "Candidate fails one configured required education/ability or mandatory unsafe-task gate.", "Eligibility fails explicitly even if other weighted components score highly."),
    ]
    rows = []
    for role_index, role in enumerate(roles, 1):
        for archetype_index, (name, scenario, expected) in enumerate(archetypes, 1):
            rows.append(case(
                f"ME-{role_index:02d}-{archetype_index}", "Matching Engine", name, "Development Team", "Critical",
                "N/A", "Deterministic API test",
                f"Published controlled opportunity exists for {role}; versioned scoring rules and weights are recorded.",
                f"{role}: {scenario}",
                "Build the controlled candidate/job input; calculate the expected score manually; call matching; compare eligibility, total percentage, component percentages and explanation; repeat once to confirm determinism.",
                expected,
            ))
    return rows


def workflow_cases():
    definitions = [
        ("WF-PUB-01", "Public", "Browse published work opportunities without signing in.", "Only public safe fields are visible."),
        ("WF-PUB-02", "Public", "Attempt to open a protected dashboard anonymously.", "The visitor is redirected to sign-in without private data."),
        ("WF-CAN-01", "Candidate", "Register, verify and sign in as a candidate.", "The verified candidate reaches the correct dashboard."),
        ("WF-CAN-02", "Candidate", "Reset and rebuild the complete candidate profile.", "Old selections are cleared and the new profile saves once."),
        ("WF-CAN-03", "Candidate", "Create, review, edit and confirm AI profile suggestions.", "Only confirmed edited items are saved."),
        ("WF-CAN-04", "Candidate", "Match and open a work opportunity.", "The result, score and explanation correspond to the saved profile."),
        ("WF-CAN-05", "Candidate", "Match and open a training opportunity.", "Training preference filters and opens the correct opportunity."),
        ("WF-CAN-06", "Candidate", "Submit a valid application with documents.", "One application is stored and acknowledged."),
        ("WF-CAN-07", "Candidate", "Track the submitted application.", "The latest authorized status is displayed."),
        ("WF-EMP-01", "Employer", "Complete or update the company profile.", "Validated company information persists correctly."),
        ("WF-EMP-02", "Employer", "Create and publish a work opportunity.", "The controlled role and tasks are published once."),
        ("WF-EMP-03", "Employer", "Create and publish a hospitality training program.", "The opportunity is identified and filtered as training."),
        ("WF-EMP-04", "Employer", "Configure education, abilities, important tasks and assistance.", "All controlled requirements persist and feed matching."),
        ("WF-EMP-05", "Employer", "Review an application and update its status.", "Only an authorized employer changes the valid application."),
        ("WF-ADM-01", "Admin", "Open and filter active and archived users.", "Authorized records load without exposing secrets."),
        ("WF-ADM-02", "Admin", "Review the controlled role/task catalogue.", "All active definitions and tasks are visible in stable order."),
        ("WF-ADM-03", "Admin", "Import or validate a catalogue dataset.", "Invalid rows are reported and valid rows are versioned safely."),
        ("WF-VER-01", "Verifier", "Open pending candidate evidence.", "Only authorized evidence is retrievable."),
        ("WF-VER-02", "Verifier", "Accept valid evidence.", "The decision, verifier and timestamp are auditable."),
        ("WF-VER-03", "Verifier", "Reject invalid evidence with a reason.", "The reason and decision are stored and visible to the candidate."),
    ]
    return [case(
        test_id, "Complete Workflows", subsystem, "Development Team", "Critical", "English baseline; repeat key UI in French/Arabic",
        "Local and staging", "Required fixture accounts and controlled catalogue data exist.", scenario,
        "Execute the complete journey from its initial state; verify each UI response and corresponding API/database state; save one representative screenshot and any failure evidence.", expected,
    ) for test_id, subsystem, scenario, expected in definitions]


def security_cases():
    definitions = [
        ("SEC-01", "Authentication", "Call a protected endpoint without a JWT.", "401; no protected data."),
        ("SEC-02", "Authentication", "Use a malformed JWT.", "401; no state change."),
        ("SEC-03", "Authentication", "Use an expired JWT.", "401 and safe sign-in recovery."),
        ("SEC-04", "Authentication", "Modify a valid JWT payload/signature.", "401; modification is detected."),
        ("SEC-05", "Authorization", "Candidate calls an employer endpoint.", "403; no employer data or mutation."),
        ("SEC-06", "Authorization", "Employer calls an administrator endpoint.", "403; no administrator data."),
        ("SEC-07", "Authorization", "Candidate requests another candidate’s object ID.", "Denied or safely scoped to the authenticated candidate."),
        ("SEC-08", "Documents", "Request a private document without authorization.", "Denied; no direct public object exposure."),
        ("SEC-09", "Input validation", "Submit unexpected enum and identifier values.", "400 with safe validation message."),
        ("SEC-10", "Input validation", "Submit SQL-injection-style text.", "Input is treated as data; query and database remain safe."),
        ("SEC-11", "Input validation", "Submit HTML/script payloads in text fields.", "No script executes when content is displayed."),
        ("SEC-12", "Uploads", "Upload a file exceeding the configured limit.", "Upload is rejected before persistence."),
        ("SEC-13", "Uploads", "Upload an unsupported MIME type.", "Upload is rejected with a safe message."),
        ("SEC-14", "Uploads", "Rename an unsupported file to an allowed extension.", "Content/MIME validation rejects it."),
        ("SEC-15", "Business logic", "Submit the same application twice.", "Duplicate submission is prevented or handled explicitly."),
        ("SEC-16", "Password reset", "Reuse a consumed password-reset token.", "Token reuse is rejected."),
        ("SEC-17", "AI consent", "Request AI suggestions without current consent/version.", "Request is rejected and narrative is not processed."),
        ("SEC-18", "AI authorization", "Confirm an invented task or job identifier.", "Unknown identifier is rejected."),
        ("SEC-19", "Voice authorization", "Request an unavailable or wrong-role voice action.", "Action is rejected without dispatch."),
        ("SEC-20", "Privacy", "Reset the candidate profile and retrieve it again.", "All intended profile selections and confirmed skills are cleared."),
    ]
    return [case(
        test_id, "Security and Privacy", subsystem, "Development Team", "Critical" if subsystem in ("Authorization", "Documents") else "High",
        "N/A", "Local first; staging safe checks only", "Use fictional test accounts and non-destructive payloads.", scenario,
        "Record the authorized baseline; issue the negative request; record HTTP status/body and database/UI state; verify no secret appears in logs or response.", expected,
    ) for test_id, subsystem, scenario, expected in definitions]


def accessibility_cases():
    journeys = [
        ("Sign in", "Reach the correct dashboard and hear/focus validation errors."),
        ("Build profile", "Complete required fields and save without an inaccessible control."),
        ("AI profile review", "Enter text, consent, review and confirm suggestions accessibly."),
        ("Match opportunity", "Generate and understand a match result and explanation."),
        ("Submit application", "Select files, handle errors and submit with confirmation."),
        ("Employer posts opportunity", "Complete controlled requirements, tasks and assistance."),
    ]
    modes = [
        ("KB", "Keyboard only", "Logical focus order, visible focus, operable controls and no keyboard trap."),
        ("ZM", "200% zoom / responsive reflow", "No two-dimensional scrolling for normal content and no hidden essential control."),
        ("SR", "JAWS or NVDA", "Labels, headings, errors, status changes and dialogs are announced meaningfully."),
    ]
    rows = []
    for journey_index, (journey, outcome) in enumerate(journeys, 1):
        for mode_code, mode, mode_expected in modes:
            rows.append(case(
                f"ACC-{journey_index:02d}-{mode_code}", "Accessibility", journey, "Development Team", "Critical",
                "English; sample Arabic RTL where relevant", mode,
                "Browser cache cleared; assistive technology/configuration recorded; test data prepared.",
                f"Complete the {journey} journey using {mode}.",
                "Start from the same initial state; complete without mouse where applicable; record blockers, assistance, focus/announcement observations and screenshots.",
                f"{outcome} {mode_expected}",
            ))
    return rows


def i18n_cases():
    screens = ["Sign in", "Candidate profile", "AI profile builder", "Job matching", "Employer opportunity creation", "Admin / verifier dashboard"]
    rows = []
    for screen_index, screen in enumerate(screens, 1):
        for code, language in LANGUAGES.items():
            rows.append(case(
                f"I18N-{screen_index:02d}-{code.upper()}", "Multilingual and RTL", screen, "Development Team", "High",
                language, "Desktop responsive baseline",
                "Translation files pass contract tests; representative controlled data exists.",
                f"Open and complete the principal interaction on {screen} in {language}.",
                "Switch language; inspect headings, controls, validation, dynamic data and direction; navigate away/back; verify entered state and language persistence; capture missing keys or overflow.",
                "No untranslated key or blocking overflow; meaning remains consistent; Arabic uses correct RTL direction while numbers and technical values remain readable.",
            ))
    return rows


def device_cases():
    journeys = ["Candidate profile", "AI suggestions", "Matching and application", "Employer opportunity creation"]
    environments = [
        ("WIN", "Windows laptop — Edge/Chrome"),
        ("AND", "Android phone — Chrome"),
        ("IOS", "iPhone/iPad — Safari"),
    ]
    rows = []
    for journey_index, journey in enumerate(journeys, 1):
        for env_code, environment in environments:
            rows.append(case(
                f"DEV-{journey_index:02d}-{env_code}", "Browser and Device", journey, "Development Team", "Critical" if journey in ("Candidate profile", "Matching and application") else "High",
                "Selected UI language recorded", environment,
                "Use deployed HTTPS staging; record device model, OS and browser version in Notes.",
                f"Complete {journey} on {environment}.",
                "Open a fresh session; complete the journey; rotate/resize where applicable; verify microphone and file picker where relevant; record completion, layout defects and screenshot.",
                "Critical controls remain visible and usable; saved data is correct; no device-specific blocker or insecure connection warning.",
            ))
    return rows


def performance_cases():
    operations = [
        ("Sign in", "Successful authenticated session."),
        ("Load candidate profile", "Profile JSON and UI load correctly."),
        ("Save candidate profile", "One validated update persists."),
        ("Generate AI suggestions", "Structured reviewable suggestions return without premature save."),
        ("Process voice command", "Authorized feedback/action returns."),
        ("Calculate matching", "Deterministic score and explanation return."),
        ("Load published opportunities", "Paginated/list data returns correctly."),
        ("Publish opportunity", "One valid opportunity persists."),
        ("Submit application", "One valid application persists without duplication."),
    ]
    loads = [
        ("L1", "1 user / normal network"),
        ("L5", "5 concurrent users or requests"),
        ("L10", "10 concurrent users or throttled network"),
    ]
    rows = []
    for operation_index, (operation, expected) in enumerate(operations, 1):
        for load_code, load in loads:
            rows.append(case(
                f"PERF-{operation_index:02d}-{load_code}", "Performance and Reliability", operation, "Development Team", "High",
                "N/A", load,
                "Staging services healthy; test data isolated; timer and server logs available.",
                f"Measure {operation} under {load}.",
                "Run the prepared operation with the stated load; record successes, failures, p50/p95 where multiple requests are used, 5xx responses and recovery behaviour; verify final data integrity.",
                f"{expected} Non-AI p95 <1 s, scoring p95 <2 s, voice p95 <10 s, AI suggestions p95 <30 s where applicable; <1% 5xx.",
            ))
    return rows


def all_cases():
    rows = []
    for builder in (
        ai_cases, transcription_cases, voice_navigation_cases, matching_cases,
        workflow_cases, security_cases, accessibility_cases, i18n_cases,
        device_cases, performance_cases,
    ):
        rows.extend(builder())
    assert len(rows) == 310, f"Expected 310 cases, got {len(rows)}"
    assert len({row[0] for row in rows}) == 310, "Test IDs must be unique"
    return rows


def title(ws, text, subtitle=None):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:K1")
    ws["A1"] = text
    ws["A1"].font = Font(size=22, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    if subtitle:
        ws.merge_cells("A2:K2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(size=11, italic=True, color=TEXT)
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 34


def style_header(row):
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def add_validation(ws, formula, column_letter, start_row, end_row):
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Choose a value from the list."
    validation.errorTitle = "Invalid value"
    validation.prompt = "Select one of the controlled values."
    validation.promptTitle = "Controlled field"
    ws.add_data_validation(validation)
    validation.add(f"{column_letter}{start_row}:{column_letter}{end_row}")


def build_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    readme = wb.create_sheet("Read Me")
    title(readme, "JoIn Final Verification Register", "Editable engineering test workbook — Development Team")
    instructions = [
        ("Purpose", "Record repeatable evidence for the final report without placing all 310 rows in the report itself."),
        ("Where to work", "Enter results only in Master Register, Defects and Evidence Index. Dashboard formulas update when Excel recalculates."),
        ("Before execution", "Freeze the tested commit/deployment URL. Do not change expected results after seeing an outcome."),
        ("First-run status", "Use Pass, Fail or Partial only after valid execution against healthy required services. Use Blocked for unavailable services, wrong paths, runner problems or unmet test-data preconditions; use Not Run when execution has not started."),
        ("After a correction", "Keep the original status. Enter the result in Retest Status and link the defect/evidence."),
        ("Percentages", "Report Passed / Executed for each area. Blocked and Not Run cases are excluded from Executed. Do not combine safety-critical failures into a misleading single score."),
        ("Evidence", "Use stable evidence IDs such as EV-AI-001 and store screenshots/logs in the project testing-evidence directory."),
        ("Privacy", "Use fictional candidates. Never place real disability evidence, passwords, API keys, JWTs or private documents in this workbook."),
        ("Ownership", "All planned and completed checks are recorded under the Development Team. Individual tester names are intentionally omitted."),
    ]
    readme["A4"] = "Field"
    readme["B4"] = "Guidance"
    style_header(readme[4][:2])
    for row_index, values in enumerate(instructions, 5):
        readme.cell(row_index, 1, values[0]).font = Font(bold=True, color=NAVY)
        readme.cell(row_index, 2, values[1])
        for column in (1, 2):
            readme.cell(row_index, column).alignment = Alignment(wrap_text=True, vertical="top")
            readme.cell(row_index, column).border = Border(bottom=THIN)
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 105
    readme.freeze_panes = "A5"
    readme["A16"] = "Workbook created"
    readme["B16"] = date.today().isoformat()
    readme["A17"] = "Planned executions"
    readme["B17"] = 310

    dashboard = wb.create_sheet("Dashboard")
    title(dashboard, "Verification Dashboard", "Summary for the report; detailed evidence remains in the Master Register")
    headers = ["Test area", "Owner", "Planned", "Executed", "First-pass passed", "Failed", "Partial / blocked", "First-pass rate", "Final passed", "Final rate", "Acceptance target"]
    for index, header in enumerate(headers, 1):
        dashboard.cell(4, index, header)
    style_header(dashboard[4])
    first_data_row = 5
    for offset, (area, planned, owner, target_text) in enumerate(AREAS):
        row = first_data_row + offset
        dashboard.cell(row, 1, area)
        dashboard.cell(row, 2, owner)
        dashboard.cell(row, 3, planned)
        dashboard.cell(row, 4, f'=COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"Pass")+COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"Fail")+COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"Partial")')
        dashboard.cell(row, 5, f'=COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"Pass")')
        dashboard.cell(row, 6, f'=COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"Fail")')
        dashboard.cell(row, 7, f'=COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"Partial")+COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"Blocked")')
        dashboard.cell(row, 8, f'=IFERROR(E{row}/D{row},0)')
        dashboard.cell(row, 9, f'=COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"Pass")+COUNTIFS(\'Master Register\'!$B:$B,$A{row},\'Master Register\'!$N:$N,"<>Pass",\'Master Register\'!$S:$S,"Pass")')
        dashboard.cell(row, 10, f'=IFERROR(I{row}/D{row},0)')
        dashboard.cell(row, 11, target_text)
        for column in range(1, 12):
            dashboard.cell(row, column).border = Border(bottom=THIN)
            dashboard.cell(row, column).alignment = Alignment(wrap_text=True, vertical="center")
        dashboard.cell(row, 8).number_format = "0.0%"
        dashboard.cell(row, 10).number_format = "0.0%"
    total_row = first_data_row + len(AREAS)
    dashboard.cell(total_row, 1, "TOTAL")
    dashboard.cell(total_row, 3, "=SUM(C5:C14)")
    dashboard.cell(total_row, 4, "=SUM(D5:D14)")
    dashboard.cell(total_row, 5, "=SUM(E5:E14)")
    dashboard.cell(total_row, 6, "=SUM(F5:F14)")
    dashboard.cell(total_row, 7, "=SUM(G5:G14)")
    dashboard.cell(total_row, 8, f"=IFERROR(E{total_row}/D{total_row},0)")
    dashboard.cell(total_row, 9, "=SUM(I5:I14)")
    dashboard.cell(total_row, 10, f"=IFERROR(I{total_row}/D{total_row},0)")
    for cell in dashboard[total_row]:
        cell.fill = PatternFill("solid", fgColor=GREY)
        cell.font = Font(bold=True, color=NAVY)
        cell.border = Border(top=THIN, bottom=THIN)
    dashboard.cell(total_row, 8).number_format = "0.0%"
    dashboard.cell(total_row, 10).number_format = "0.0%"
    for letter, width in {"A": 30, "B": 13, "C": 11, "D": 11, "E": 16, "F": 10, "G": 16, "H": 14, "I": 12, "J": 12, "K": 48}.items():
        dashboard.column_dimensions[letter].width = width
    dashboard.freeze_panes = "A5"
    dashboard.conditional_formatting.add(f"J5:J{total_row}", FormulaRule(formula=["J5>=0.95"], fill=PatternFill("solid", fgColor=LIGHT_GREEN)))
    dashboard.conditional_formatting.add(f"J5:J{total_row}", FormulaRule(formula=["AND(J5>0,J5<0.95)"], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    chart = BarChart()
    chart.title = "Final pass rate by test area"
    chart.y_axis.title = "Pass rate"
    chart.x_axis.title = "Test area"
    chart.height = 8
    chart.width = 18
    chart.add_data(Reference(dashboard, min_col=10, min_row=4, max_row=14), titles_from_data=True)
    chart.set_categories(Reference(dashboard, min_col=1, min_row=5, max_row=14))
    dashboard.add_chart(chart, "A18")

    master = wb.create_sheet("Master Register")
    headers = [
        "Test ID", "Test Area", "Subsystem / Group", "Owner", "Priority", "Language",
        "Condition / Environment", "Preconditions", "Test Input / Scenario", "Execution Steps",
        "Expected Result", "Actual Result", "Evidence Reference", "First-run Status",
        "Severity if Failed", "Duration (seconds)", "Response Time (ms)", "Run Date",
        "Retest Status", "Retest Date", "Notes",
    ]
    master.append(headers)
    for row in all_cases():
        master.append(row)
    style_header(master[1])
    master.freeze_panes = "A2"
    master.auto_filter.ref = f"A1:U{master.max_row}"
    master.sheet_view.showGridLines = False
    widths = [16, 29, 28, 12, 12, 16, 28, 42, 52, 58, 52, 46, 24, 17, 18, 18, 19, 15, 16, 15, 38]
    for index, width in enumerate(widths, 1):
        master.column_dimensions[master.cell(1, index).column_letter].width = width
    for row in master.iter_rows(min_row=2, max_row=master.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
        row[0].font = Font(bold=True, color=NAVY)
        row[15].number_format = "0.0"
        row[16].number_format = "0"
        row[17].number_format = "yyyy-mm-dd"
        row[19].number_format = "yyyy-mm-dd"
    table = Table(displayName="MasterTestRegister", ref=f"A1:U{master.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    master.add_table(table)

    lists = wb.create_sheet("Lists")
    list_values = {
        "A": ["Status", "Pass", "Fail", "Partial", "Blocked", "Not Run"],
        "B": ["Retest", "Pass", "Fail", "Blocked", "Not Retested"],
        "C": ["Severity", "Critical", "High", "Medium", "Low", "N/A"],
        "D": ["Priority", "Critical", "High", "Medium", "Low"],
        "E": ["Owner", "Development Team"],
        "F": ["Defect status", "Open", "In Progress", "Ready for Retest", "Closed", "Accepted Risk"],
        "G": ["Evidence type", "Screenshot", "Video", "Terminal log", "JSON report", "CSV", "Accessibility notes", "Other"],
    }
    for column, values in list_values.items():
        for row_index, value in enumerate(values, 1):
            lists[f"{column}{row_index}"] = value
    lists.sheet_state = "hidden"
    add_validation(master, "=Lists!$A$2:$A$6", "N", 2, 2000)
    add_validation(master, "=Lists!$B$2:$B$5", "S", 2, 2000)
    add_validation(master, "=Lists!$C$2:$C$6", "O", 2, 2000)
    add_validation(master, "=Lists!$D$2:$D$5", "E", 2, 2000)
    add_validation(master, "=Lists!$E$2:$E$4", "D", 2, 2000)
    master.conditional_formatting.add(f"N2:N{master.max_row}", FormulaRule(formula=['N2="Pass"'], fill=PatternFill("solid", fgColor=LIGHT_GREEN)))
    master.conditional_formatting.add(f"N2:N{master.max_row}", FormulaRule(formula=['N2="Fail"'], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    master.conditional_formatting.add(f"N2:N{master.max_row}", FormulaRule(formula=['OR(N2="Partial",N2="Blocked")'], fill=PatternFill("solid", fgColor=LIGHT_ORANGE)))
    master.conditional_formatting.add(f"S2:S{master.max_row}", FormulaRule(formula=['S2="Pass"'], fill=PatternFill("solid", fgColor=LIGHT_GREEN)))

    metrics = wb.create_sheet("Metrics")
    title(metrics, "Accuracy and Performance Metrics", "Enter aggregate counts after completing the corresponding Master Register rows")
    metrics["A4"] = "AI PROFILE METRICS"
    metrics["A4"].font = Font(bold=True, color=ORANGE, size=13)
    ai_headers = ["Language", "Text quality", "TP", "FP", "FN", "Precision", "Recall", "F1", "False disability inference", "Invented task"]
    for index, header in enumerate(ai_headers, 1):
        metrics.cell(5, index, header)
    style_header(metrics[5][:10])
    row_index = 6
    for language in LANGUAGES.values():
        for quality in ("Clean text", "Imperfect transcript"):
            metrics.cell(row_index, 1, language)
            metrics.cell(row_index, 2, quality)
            metrics.cell(row_index, 6, f"=IFERROR(C{row_index}/(C{row_index}+D{row_index}),0)")
            metrics.cell(row_index, 7, f"=IFERROR(C{row_index}/(C{row_index}+E{row_index}),0)")
            metrics.cell(row_index, 8, f"=IFERROR(2*F{row_index}*G{row_index}/(F{row_index}+G{row_index}),0)")
            for column in (6, 7, 8):
                metrics.cell(row_index, column).number_format = "0.0%"
            row_index += 1

    metrics["A14"] = "VOICE METRICS"
    metrics["A14"].font = Font(bold=True, color=BLUE, size=13)
    voice_headers = ["Language", "Request type", "Executed", "Transcript correct", "Classification correct", "Final action correct", "Unauthorized actions", "Average ms", "p95 ms", "Notes"]
    for index, header in enumerate(voice_headers, 1):
        metrics.cell(15, index, header)
    style_header(metrics[15][:10])
    row_index = 16
    for language in LANGUAGES.values():
        for request_type in ("Navigation", "Interface action", "Page question", "Confirmation / rejected"):
            metrics.cell(row_index, 1, language)
            metrics.cell(row_index, 2, request_type)
            row_index += 1

    metrics["A30"] = "PERFORMANCE METRICS"
    metrics["A30"].font = Font(bold=True, color=GREEN, size=13)
    perf_headers = ["Operation", "Load", "Runs", "Successful", "Failed", "Success rate", "p50 ms", "p95 ms", "5xx", "Notes"]
    for index, header in enumerate(perf_headers, 1):
        metrics.cell(31, index, header)
    style_header(metrics[31][:10])
    row_index = 32
    operations = ["Sign in", "Load profile", "Save profile", "AI suggestions", "Voice command", "Matching", "Load opportunities", "Publish opportunity", "Submit application"]
    for operation in operations:
        for load in ("1 user", "5 concurrent", "10 concurrent / throttled"):
            metrics.cell(row_index, 1, operation)
            metrics.cell(row_index, 2, load)
            metrics.cell(row_index, 6, f"=IFERROR(D{row_index}/C{row_index},0)")
            metrics.cell(row_index, 6).number_format = "0.0%"
            row_index += 1
    for column in range(1, 11):
        metrics.column_dimensions[get_column_letter(column)].width = 24 if column in (1, 2, 10) else 18
    for row in metrics.iter_rows(min_row=5, max_row=metrics.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
    metrics.freeze_panes = "A5"

    defects = wb.create_sheet("Defects")
    defect_headers = ["Defect ID", "Related Test ID", "Area", "Summary", "Severity", "Reproduction", "Expected", "Actual", "Root Cause", "Correction", "Owner", "Status", "Opened Date", "Retest Test ID", "Closed Date", "Evidence Reference"]
    defects.append(defect_headers)
    for index in range(1, 31):
        defects.append([f"DEF-{index:03d}"] + [""] * (len(defect_headers) - 1))
    style_header(defects[1])
    defects.freeze_panes = "A2"
    defects.auto_filter.ref = f"A1:P{defects.max_row}"
    defects.sheet_view.showGridLines = False
    for column in range(1, len(defect_headers) + 1):
        defects.column_dimensions[defects.cell(1, column).column_letter].width = 18 if column not in (4, 6, 7, 8, 9, 10) else 38
    for row in defects.iter_rows(min_row=2, max_row=defects.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
        row[0].font = Font(bold=True, color=NAVY)
        row[12].number_format = "yyyy-mm-dd"
        row[14].number_format = "yyyy-mm-dd"
    add_validation(defects, "=Lists!$C$2:$C$6", "E", 2, 1000)
    add_validation(defects, "=Lists!$E$2:$E$4", "K", 2, 1000)
    add_validation(defects, "=Lists!$F$2:$F$6", "L", 2, 1000)
    defect_table = Table(displayName="DefectRegister", ref=f"A1:P{defects.max_row}")
    defect_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    defects.add_table(defect_table)

    evidence = wb.create_sheet("Evidence Index")
    evidence_headers = ["Evidence ID", "Related Test ID(s)", "Evidence Type", "File Path / URL", "Description", "Captured Date", "Owner", "Verified", "Notes"]
    evidence.append(evidence_headers)
    for index in range(1, 51):
        evidence.append([f"EV-{index:03d}"] + [""] * (len(evidence_headers) - 1))
    style_header(evidence[1])
    evidence.freeze_panes = "A2"
    evidence.auto_filter.ref = f"A1:I{evidence.max_row}"
    evidence.sheet_view.showGridLines = False
    for column, width in enumerate([16, 24, 20, 55, 48, 16, 14, 14, 34], 1):
        evidence.column_dimensions[evidence.cell(1, column).column_letter].width = width
    for row in evidence.iter_rows(min_row=2, max_row=evidence.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
        row[0].font = Font(bold=True, color=NAVY)
        row[5].number_format = "yyyy-mm-dd"
    add_validation(evidence, "=Lists!$G$2:$G$8", "C", 2, 1000)
    add_validation(evidence, "=Lists!$E$2:$E$4", "G", 2, 1000)
    verified_validation = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    evidence.add_data_validation(verified_validation)
    verified_validation.add("H2:H1000")
    evidence_table = Table(displayName="EvidenceRegister", ref=f"A1:I{evidence.max_row}")
    evidence_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    evidence.add_table(evidence_table)

    wb.move_sheet(lists, offset=len(wb.sheetnames))
    return wb


def verify_workbook(path):
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["Read Me", "Dashboard", "Master Register", "Metrics", "Defects", "Evidence Index", "Lists"]
    assert workbook["Master Register"].max_row == 311
    assert workbook["Master Register"]["A311"].value == "PERF-09-L10"
    assert workbook["Dashboard"]["C15"].value == "=SUM(C5:C14)"
    assert len(workbook["Master Register"].tables) == 1


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook()
    workbook.save(OUTPUT)
    verify_workbook(OUTPUT)
    print(f"Created and verified: {OUTPUT.resolve()}")
    print("Planned test executions: 310")
