from __future__ import annotations

import sys
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(sys.argv[1] if len(sys.argv) > 1 else "JoIn_Final_Test_Register.xlsx")
REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
MATCHING_REPORT = Path(sys.argv[2]) if len(sys.argv) > 2 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "matching-quality-results.json"
I18N_REPORT = Path(sys.argv[3]) if len(sys.argv) > 3 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "i18n-quality-results.json"
LOAD_REPORT = Path(sys.argv[4]) if len(sys.argv) > 4 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "safe-load-results.json"
PERFORMANCE_REPORT = Path(sys.argv[5]) if len(sys.argv) > 5 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "performance-results.json"
STATEFUL_REPORT = Path(sys.argv[6]) if len(sys.argv) > 6 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "stateful-load-results.json"
SECURITY_REPORT = Path(sys.argv[7]) if len(sys.argv) > 7 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "security-quality-results.json"
WCAG_REPORT = Path(sys.argv[8]) if len(sys.argv) > 8 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "wcag-deployed-results.json"
WORKFLOW_REPORT = Path(sys.argv[9]) if len(sys.argv) > 9 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "workflow-quality-results.json"
AI_REPORT = Path(sys.argv[10]) if len(sys.argv) > 10 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "profile-ai" / "ai-quality-matrix-results.json"
VERIFIER_REPORT = Path(sys.argv[11]) if len(sys.argv) > 11 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "verifier-workflow-results.json"
AI_PERFORMANCE_REPORT = Path(sys.argv[12]) if len(sys.argv) > 12 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "ai-performance-results.json"
MUTATION_PERFORMANCE_REPORT = Path(sys.argv[13]) if len(sys.argv) > 13 else REPOSITORY_ROOT / "docs" / "testing" / "staging-2026-09-10" / "mutation-performance-results.json"


RESULTS = {
    "I18N-01-EN": {
        "actual": "The English sign-in page rendered with a semantic main region, visible heading and no page error.",
        "evidence": "EV-I18N-001",
        "status": "Pass",
        "notes": "Evaluated output case. Render services were healthy before execution.",
    },
    "I18N-01-FR": {
        "actual": "French switching updated the document language, retained LTR direction, translated the sign-in heading and title, synchronized the voice-language control and persisted after reload.",
        "evidence": "EV-I18N-001",
        "status": "Pass",
        "notes": "Evaluated output case. Render services were healthy before execution.",
    },
    "I18N-01-AR": {
        "actual": "Arabic switching updated lang and RTL direction; sign-in rendered without horizontal overflow, kept the email input LTR and synchronized the voice-language control.",
        "evidence": "EV-I18N-001",
        "status": "Pass",
        "notes": "Evaluated output case. Render services were healthy before execution.",
    },
}


def update_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    master = workbook["Master Register"]
    rows_by_id = {master.cell(row, 1).value: row for row in range(2, master.max_row + 1)}
    for test_id, result in RESULTS.items():
        row = rows_by_id[test_id]
        master.cell(row, 12, result["actual"])
        master.cell(row, 13, result["evidence"])
        master.cell(row, 14, result["status"])
        master.cell(row, 18, date(2026, 9, 10))
        master.cell(row, 21, result["notes"])

    matching_report = json.loads(MATCHING_REPORT.read_text(encoding="utf-8"))
    i18n_report = json.loads(I18N_REPORT.read_text(encoding="utf-8"))
    load_report = json.loads(LOAD_REPORT.read_text(encoding="utf-8"))
    performance_report = json.loads(PERFORMANCE_REPORT.read_text(encoding="utf-8"))
    stateful_report = json.loads(STATEFUL_REPORT.read_text(encoding="utf-8"))
    security_report = json.loads(SECURITY_REPORT.read_text(encoding="utf-8"))
    wcag_report = json.loads(WCAG_REPORT.read_text(encoding="utf-8"))
    workflow_report = json.loads(WORKFLOW_REPORT.read_text(encoding="utf-8"))
    ai_report = json.loads(AI_REPORT.read_text(encoding="utf-8"))
    verifier_report = json.loads(VERIFIER_REPORT.read_text(encoding="utf-8"))
    ai_performance_report = json.loads(AI_PERFORMANCE_REPORT.read_text(encoding="utf-8"))
    mutation_performance_report = json.loads(MUTATION_PERFORMANCE_REPORT.read_text(encoding="utf-8"))
    ai_cases_by_id = {item["testId"]: item for item in ai_report["cases"]}
    workflow_cases_by_id = {item["testId"]: item for item in workflow_report["results"]}
    workflow_cases_by_id.update({item["testId"]: item for item in verifier_report["results"]})
    matching_cases = {
        (item["role"], item["archetype"]): item
        for item in matching_report["cases"]
    }
    matching_roles = sorted({item["role"] for item in matching_report["cases"]}, key=len, reverse=True)
    archetype_names = {
        "Independent": "Independent",
        "With support": "With support",
        "Ineligible gate": "Eligibility gates",
    }
    matching_rows = 0
    ai_rows = 0
    performance_rows = 0
    for row in range(2, master.max_row + 1):
        area = master.cell(row, 2).value
        if area == "Matching Engine":
            scenario = str(master.cell(row, 9).value)
            role = next(name for name in matching_roles if scenario.startswith(f"{name}:"))
            case = matching_cases[(role, archetype_names[master.cell(row, 3).value])]
            status = "Pass" if case["status"] == "passed" else "Blocked"
            if status == "Pass":
                actual = (
                    f"Validated {case['offerCount']} deployed offer(s): {case['eligibleCount']} eligible and "
                    f"{case['ineligibleCount']} ineligible. Formula components, feasibility factors, explanations "
                    "and the repeated deterministic response matched the independent oracle."
                )
            else:
                actual = case["note"]
            master.cell(row, 12, actual)
            master.cell(row, 13, "EV-MATCH-001")
            master.cell(row, 14, status)
            master.cell(row, 16, round(float(case.get("durationMs", 0)) / 1000, 1) if case.get("durationMs") else None)
            master.cell(row, 18, date(2026, 9, 10))
            master.cell(row, 21, "Blocked cases are excluded from the executed success rate." if status == "Blocked" else "Evaluated output case against healthy API and scoring services.")
            matching_rows += 1
        elif area == "AI Profile Suggestions":
            test_id = master.cell(row, 1).value
            item = ai_cases_by_id[test_id]
            status = item["status"]
            if status == "Pass":
                actual = f"{item['actual']} HTTP {item.get('httpStatus')}; the profile remained unchanged before cleanup={item.get('profileUnchangedBeforeCleanup')}."
            elif status == "Fail":
                failed_checks = [check["name"] for check in item.get("evaluation", {}).get("checks", []) if not check.get("passed")]
                actual = f"{item['actual']} Missing or incorrect expectation: {', '.join(failed_checks)}. HTTP {item.get('httpStatus')}."
            else:
                actual = item.get("actual", "The output could not be evaluated.")
            master.cell(row, 12, actual)
            master.cell(row, 13, "EV-AI-002")
            master.cell(row, 14, status)
            master.cell(row, 16, round(float(item.get("durationMs", 0)) / 1000, 1) if item.get("durationMs") else None)
            master.cell(row, 17, round(float(item.get("durationMs", 0))) if item.get("durationMs") else None)
            master.cell(row, 18, date(2026, 9, 10))
            master.cell(row, 21, "Evaluated structured AI output against the gold expectation; suggestions were not saved and the synthetic profile was reset." if status != "Blocked" else "Excluded from the executed success rate.")
            ai_rows += 1
    i18n_cases = {item["testId"]: item for item in i18n_report["cases"]}
    for test_id, item in i18n_cases.items():
        row = rows_by_id[test_id]
        observation = item.get("observation", {})
        if item["status"] == "passed":
            actual = (
                f"Rendered with lang={observation.get('lang')}, dir={observation.get('dir')}, a visible localized heading, "
                "no unresolved translation key, no page error and no blocking horizontal overflow."
            )
            status = "Pass"
        else:
            actual = item.get("reason", "The multilingual output did not meet the expected result.")
            status = "Fail"
        master.cell(row, 12, actual)
        master.cell(row, 13, "EV-I18N-002")
        master.cell(row, 14, status)
        master.cell(row, 16, round(float(item.get("durationMs", 0)) / 1000, 1) if item.get("durationMs") else None)
        master.cell(row, 18, date(2026, 9, 10))
        master.cell(row, 21, "Evaluated deployed multilingual/RTL output case.")

    operation_numbers = {"Sign in": "01", "Load profile": "02", "Load opportunities": "07"}
    load_codes = {1: "L1", 5: "L5", 10: "L10"}
    baseline_jobs_p95 = next(item["p95Ms"] for item in performance_report["api"] if item["name"] == "public jobs")
    for item in load_report["cases"]:
        test_id = f"PERF-{operation_numbers[item['operation']]}-{load_codes[int(item['load'])]}"
        row = rows_by_id[test_id]
        status = "Pass" if item["status"] == "passed" else "Fail"
        threshold = item["thresholdMs"]
        extra = ""
        if item["operation"] == "Load opportunities" and int(item["load"]) == 1:
            status = "Fail"
            threshold = performance_report["thresholds"]["apiP95Ms"]
            extra = f" A separate five-sample baseline measured p95={baseline_jobs_p95} ms against the project API target."
        actual = (
            f"{item['successful']}/{item['requests']} requests returned successfully with 0 server errors; "
            f"p50={item['p50Ms']} ms and p95={item['p95Ms']} ms versus the {threshold} ms target.{extra}"
        )
        master.cell(row, 12, actual)
        master.cell(row, 13, "EV-PERF-001, EV-PERF-002")
        master.cell(row, 14, status)
        master.cell(row, 17, item["p95Ms"])
        master.cell(row, 18, date(2026, 9, 10))
        master.cell(row, 21, "Correct HTTP outputs were returned; status reflects latency only.")
        performance_rows += 1
    stateful_operation_numbers = {"Save profile": "03", "Matching": "06"}
    for item in stateful_report["cases"]:
        test_id = f"PERF-{stateful_operation_numbers[item['operation']]}-{load_codes[int(item['load'])]}"
        row = rows_by_id[test_id]
        status = "Pass" if item["status"] == "passed" else "Fail"
        actual = (
            f"{item['successful']}/{item['requests']} requests returned successfully, with {item['serverErrors']} server errors; "
            f"p50={item['p50Ms']} ms and p95={item['p95Ms']} ms versus the {item['thresholdMs']} ms target."
        )
        master.cell(row, 12, actual)
        master.cell(row, 13, "EV-PERF-003")
        master.cell(row, 14, status)
        master.cell(row, 17, item["p95Ms"])
        master.cell(row, 18, date(2026, 9, 10))
        master.cell(row, 21, "Synthetic candidate data was reset after the test. Failures reflect observed 5xx responses or threshold overruns.")
        performance_rows += 1
    for item in ai_performance_report["cases"]:
        load_code = load_codes[int(item["load"])]
        row = rows_by_id[f"PERF-04-{load_code}"]
        all_dependency_unavailable = item["successful"] == 0 and item.get("httpStatuses") == {"503": item["requests"]}
        status = "Blocked" if all_dependency_unavailable else ("Pass" if item["status"] == "passed" else "Fail")
        actual = (
            f"{item['successful']}/{item['requests']} requests returned structured suggestions; "
            f"HTTP statuses={item.get('httpStatuses')}; p50={item['p50Ms']} ms and p95={item['p95Ms']} ms."
        )
        if all_dependency_unavailable:
            actual += " A post-run probe confirmed API health HTTP 200 but the upstream AI path still returned HTTP 503, so inference latency could not be evaluated."
        master.cell(row, 12, actual)
        master.cell(row, 13, "EV-PERF-004")
        master.cell(row, 14, status)
        master.cell(row, 17, item["p95Ms"])
        master.cell(row, 18, date(2026, 9, 10))
        master.cell(row, 21, "Excluded from executed performance rates because the deployed upstream AI dependency was unavailable during and after the run." if status == "Blocked" else "Evaluated against the 30-second AI p95 target.")
    mutation_operation_numbers = {"Publish opportunity": "08", "Submit application": "09"}
    for item in mutation_performance_report["cases"]:
        test_id = f"PERF-{mutation_operation_numbers[item['operation']]}-{load_codes[int(item['load'])]}"
        row = rows_by_id[test_id]
        status = "Pass" if item["status"] == "passed" else "Fail"
        actual = (
            f"{item['successful']}/{item['requests']} isolated mutations returned the expected HTTP 201 with {item['serverErrors']} server errors; "
            f"p50={item['p50Ms']} ms and p95={item['p95Ms']} ms versus the {item['thresholdMs']} ms target."
        )
        master.cell(row, 12, actual)
        master.cell(row, 13, "EV-PERF-005")
        master.cell(row, 14, status)
        master.cell(row, 17, item["p95Ms"])
        master.cell(row, 18, date(2026, 9, 10))
        master.cell(row, 21, "All functional outputs succeeded; Fail indicates only that the strict latency target was exceeded. All created records were deleted.")
    security_rows = 0
    for item in security_report["results"]:
        row = rows_by_id[item["testId"]]
        status = item["status"]
        master.cell(row, 12, item["actual"])
        master.cell(row, 13, "EV-SEC-002")
        master.cell(row, 14, status)
        master.cell(row, 18, date(2026, 9, 10))
        if status == "Blocked":
            notes = "Excluded from the executed security rate because the required staging precondition was unavailable."
        elif status == "Fail":
            notes = "Evaluated against a healthy staging service; this is a genuine observed security-control failure. Temporary data was removed."
        else:
            notes = "Evaluated negative security/privacy output case against healthy staging services using synthetic data."
        master.cell(row, 21, notes)
        security_rows += 1
    workflow_rows = 0
    for test_id, item in workflow_cases_by_id.items():
        if test_id not in rows_by_id:
            continue
        row = rows_by_id[test_id]
        status = item["status"]
        is_verifier = test_id.startswith("WF-VER-")
        master.cell(row, 12, item["actual"])
        master.cell(row, 13, "EV-WF-002" if is_verifier else "EV-WF-001")
        master.cell(row, 14, status)
        master.cell(row, 18, date(2026, 9, 10))
        if status == "Blocked":
            notes = "Excluded from the executed workflow rate because an isolated staging precondition is unavailable."
        else:
            notes = "Evaluated deployed workflow output with synthetic data; state-changing records were removed after execution."
        master.cell(row, 21, notes)
        workflow_rows += 1
    assert matching_rows == 45
    assert ai_rows == 60
    assert len(i18n_cases) == 18
    assert performance_rows == 15
    assert security_rows == 20
    assert len(ai_cases_by_id) == 60
    assert workflow_rows == 20

    evidence = workbook["Evidence Index"]
    evidence.cell(2, 1, "EV-I18N-001")
    evidence.cell(2, 2, "I18N-01-EN, I18N-01-FR, I18N-01-AR")
    evidence.cell(2, 3, "JSON report")
    evidence.cell(2, 4, "docs/testing/staging-2026-09-10/i18n-results.json")
    evidence.cell(2, 5, "Translation contract and deployed language, direction, persistence, reflow and voice-language synchronization checks.")
    evidence.cell(2, 6, date(2026, 9, 10))
    evidence.cell(2, 7, "Development Team")
    evidence.cell(2, 8, "Yes")
    evidence.cell(2, 9, "Only the three sign-in language rows are scored in this first batch.")

    evidence.cell(3, 1, "EV-MATCH-001")
    evidence.cell(3, 2, "ME-01-1 through ME-15-3")
    evidence.cell(3, 3, "JSON report")
    evidence.cell(3, 4, "docs/testing/staging-2026-09-10/matching-quality-results.json")
    evidence.cell(3, 5, "45 deployed matching cases across 15 roles and three synthetic candidate conditions; independent formula oracle and repeated determinism check.")
    evidence.cell(3, 6, date(2026, 9, 10))
    evidence.cell(3, 7, "Development Team")
    evidence.cell(3, 8, "Yes")
    evidence.cell(3, 9, "45/45 passed after three disposable requirement/support fixtures supplied the missing scenario preconditions; all fixtures were deleted afterward.")

    evidence.cell(4, 1, "EV-AI-001")
    evidence.cell(4, 2, "Historical AI service recovery evidence")
    evidence.cell(4, 3, "JSON report")
    evidence.cell(4, 4, "docs/testing/staging-2026-09-10/profile-ai/ai-service-probe.json")
    evidence.cell(4, 5, "Historical AI availability probe retained to document the transient 503 condition before the service recovered.")
    evidence.cell(4, 6, date(2026, 9, 10))
    evidence.cell(4, 7, "Development Team")
    evidence.cell(4, 8, "Yes")
    evidence.cell(4, 9, "Superseded by EV-AI-002 after a successful service recovery check and complete 60-case execution.")

    evidence.cell(5, 1, "EV-SEC-001")
    evidence.cell(5, 2, "Supporting prerequisite evidence")
    evidence.cell(5, 3, "Terminal log")
    evidence.cell(5, 4, "docs/testing/staging-2026-09-10/security-route-audit.txt")
    evidence.cell(5, 5, "51 anonymous, forged-token, cross-role, valid-role and public-route authorization checks against the deployed API.")
    evidence.cell(5, 6, date(2026, 9, 10))
    evidence.cell(5, 7, "Development Team")
    evidence.cell(5, 8, "Yes")
    evidence.cell(5, 9, "51/51 passed. Kept separate from output-quality percentages as requested.")

    evidence.cell(6, 1, "EV-I18N-002")
    evidence.cell(6, 2, "I18N-01-EN through I18N-06-AR")
    evidence.cell(6, 3, "JSON report")
    evidence.cell(6, 4, "docs/testing/staging-2026-09-10/i18n-quality-results.json")
    evidence.cell(6, 5, "18 deployed multilingual/RTL output cases across six screens and three interface languages.")
    evidence.cell(6, 6, date(2026, 9, 10))
    evidence.cell(6, 7, "Development Team")
    evidence.cell(6, 8, "Yes")
    evidence.cell(6, 9, "18/18 passed.")

    evidence.cell(7, 1, "EV-PERF-001")
    evidence.cell(7, 2, "PERF-01, PERF-02 and PERF-07 groups")
    evidence.cell(7, 3, "JSON report")
    evidence.cell(7, 4, "docs/testing/staging-2026-09-10/safe-load-results.json")
    evidence.cell(7, 5, "Non-destructive API measurements at concurrency levels 1, 5 and 10; three rounds per level.")
    evidence.cell(7, 6, date(2026, 9, 10))
    evidence.cell(7, 7, "Development Team")
    evidence.cell(7, 8, "Yes")
    evidence.cell(7, 9, "All requests returned correctly with zero 5xx responses; several cases exceeded the latency thresholds.")

    evidence.cell(8, 1, "EV-PERF-002")
    evidence.cell(8, 2, "PERF-07-L1 and deployed page/API baseline")
    evidence.cell(8, 3, "JSON report")
    evidence.cell(8, 4, "docs/testing/staging-2026-09-10/performance-results.json")
    evidence.cell(8, 5, "Five-sample deployed API and page timing baseline plus production bundle-size checks.")
    evidence.cell(8, 6, date(2026, 9, 10))
    evidence.cell(8, 7, "Development Team")
    evidence.cell(8, 8, "Yes")
    evidence.cell(8, 9, "Public jobs API p95 exceeded the 1000 ms project target; other measured APIs, pages and bundles met their budgets.")

    evidence.cell(9, 1, "EV-PERF-003")
    evidence.cell(9, 2, "PERF-03-L1 through PERF-03-L10; PERF-06-L1 through PERF-06-L10")
    evidence.cell(9, 3, "JSON report")
    evidence.cell(9, 4, "docs/testing/staging-2026-09-10/stateful-load-results.json")
    evidence.cell(9, 5, "Idempotent candidate-profile save and deterministic matching measurements at concurrency levels 1, 5 and 10.")
    evidence.cell(9, 6, date(2026, 9, 10))
    evidence.cell(9, 7, "Development Team")
    evidence.cell(9, 8, "Yes")
    evidence.cell(9, 9, "Synthetic candidate data was reset. Concurrent saves produced 5xx errors at loads 5 and 10; matching load 10 exceeded its latency threshold.")

    evidence.cell(10, 1, "EV-SEC-002")
    evidence.cell(10, 2, "SEC-01 through SEC-20")
    evidence.cell(10, 3, "JSON report")
    evidence.cell(10, 4, "docs/testing/staging-2026-09-10/security-quality-results.json")
    evidence.cell(10, 5, "Twenty deployed negative security/privacy scenarios covering JWT validation, role isolation, object/document access, input and upload validation, duplicate applications, consent, controlled IDs, voice authorization and profile reset.")
    evidence.cell(10, 6, date(2026, 9, 10))
    evidence.cell(10, 7, "Development Team")
    evidence.cell(10, 8, "Yes")
    evidence.cell(10, 9, f"{security_report['summary']['pass']} passed, {security_report['summary']['fail']} failed and {security_report['summary']['blocked']} blocked. Synthetic application/profile data was cleaned after execution.")

    wcag_results = wcag_report["results"]
    scan_errors = sum(1 for item in wcag_results if item.get("scanError"))
    violation_scenarios = sum(1 for item in wcag_results if item.get("axe"))
    overflow_scenarios = sum(1 for item in wcag_results if item.get("layout", {}).get("horizontalOverflow"))
    evidence.cell(11, 1, "EV-ACC-001")
    evidence.cell(11, 2, "Supporting accessibility evidence")
    evidence.cell(11, 3, "JSON report")
    evidence.cell(11, 4, "docs/testing/staging-2026-09-10/wcag-deployed-results.json")
    evidence.cell(11, 5, "Automated WCAG 2.1 A/AA, responsive reflow, text-spacing, keyboard-tab and browser-error evidence across 13 deployed routes and eight layouts, with 104 screenshots.")
    evidence.cell(11, 6, date(2026, 9, 10))
    evidence.cell(11, 7, "Development Team")
    evidence.cell(11, 8, "Yes")
    evidence.cell(11, 9, f"104 scenarios; {scan_errors} scan errors, {violation_scenarios} scenarios with axe findings and {overflow_scenarios} with horizontal overflow. Automated evidence does not replace manual journey or screen-reader testing.")

    evidence.cell(12, 1, "EV-AI-002")
    evidence.cell(12, 2, "AI-EN-C01 through AI-AR-I10")
    evidence.cell(12, 3, "JSON report")
    evidence.cell(12, 4, "docs/testing/staging-2026-09-10/profile-ai/ai-quality-matrix-results.json")
    evidence.cell(12, 5, "Sixty gold-standard clean and imperfect profile narratives across English, French and Arabic, evaluated against structured field, ability, preference, disability, position, task and safety expectations.")
    evidence.cell(12, 6, date(2026, 9, 10))
    evidence.cell(12, 7, "Development Team")
    evidence.cell(12, 8, "Yes")
    evidence.cell(12, 9, f"{ai_report['summary']['pass']}/60 passed, {ai_report['summary']['fail']} failed and {ai_report['summary']['blocked']} blocked; no suggestion was persisted during the matrix.")

    evidence.cell(13, 1, "EV-AI-003")
    evidence.cell(13, 2, "WF-CAN-03 representative browser confirmation")
    evidence.cell(13, 3, "Screenshots and JSON report")
    evidence.cell(13, 4, "docs/testing/staging-2026-09-10/profile-ai/screenshots-en-confirmed/")
    evidence.cell(13, 5, "Representative deployed browser journey showing consent, editable structured suggestions, confirmation and saved-profile verification.")
    evidence.cell(13, 6, date(2026, 9, 10))
    evidence.cell(13, 7, "Development Team")
    evidence.cell(13, 8, "Yes")
    evidence.cell(13, 9, "The synthetic candidate profile was reset after verification.")

    evidence.cell(14, 1, "EV-WF-001")
    evidence.cell(14, 2, "WF-PUB-01 through WF-ADM-03")
    evidence.cell(14, 3, "JSON reports")
    evidence.cell(14, 4, "docs/testing/staging-2026-09-10/workflow-quality-results.json")
    evidence.cell(14, 5, "Deployed public, candidate, employer and administrator workflow outputs, supported by prepared integration and browser E2E reports.")
    evidence.cell(14, 6, date(2026, 9, 10))
    evidence.cell(14, 7, "Development Team")
    evidence.cell(14, 8, "Yes")
    evidence.cell(14, 9, "Synthetic profile, applications and opportunities were removed after execution; isolated email verification and catalogue replacement remain blocked.")

    evidence.cell(15, 1, "EV-WF-002")
    evidence.cell(15, 2, "WF-VER-01 through WF-VER-03")
    evidence.cell(15, 3, "JSON report")
    evidence.cell(15, 4, "docs/testing/staging-2026-09-10/verifier-workflow-results.json")
    evidence.cell(15, 5, "Deployed verifier document retrieval, approval and rejection using two disposable candidate accounts and synthetic private PDF evidence.")
    evidence.cell(15, 6, date(2026, 9, 10))
    evidence.cell(15, 7, "Development Team")
    evidence.cell(15, 8, "Yes")
    evidence.cell(15, 9, "3/3 passed; both disposable accounts and private documents were deleted afterward.")

    evidence.cell(16, 1, "EV-PERF-004")
    evidence.cell(16, 2, "PERF-04-L1 through PERF-04-L10")
    evidence.cell(16, 3, "JSON reports")
    evidence.cell(16, 4, "docs/testing/staging-2026-09-10/ai-performance-results.json")
    evidence.cell(16, 5, "Deployed AI suggestion load attempt at concurrency 1, 5 and 10 plus a post-run dependency probe.")
    evidence.cell(16, 6, date(2026, 9, 10))
    evidence.cell(16, 7, "Development Team")
    evidence.cell(16, 8, "Yes")
    evidence.cell(16, 9, "All 48 attempts returned HTTP 503; a post-run probe confirmed the same condition while API health remained 200. These three cases remain Blocked rather than Fail because no inference output was available to time.")

    evidence.cell(17, 1, "EV-PERF-005")
    evidence.cell(17, 2, "PERF-08-L1 through PERF-09-L10")
    evidence.cell(17, 3, "JSON report")
    evidence.cell(17, 4, "docs/testing/staging-2026-09-10/mutation-performance-results.json")
    evidence.cell(17, 5, "Opportunity publishing and application submission at concurrency 1, 5 and 10, with three rounds and a unique disposable record per request.")
    evidence.cell(17, 6, date(2026, 9, 10))
    evidence.cell(17, 7, "Development Team")
    evidence.cell(17, 8, "Yes")
    evidence.cell(17, 9, "96/96 mutations returned the correct HTTP 201 output with zero 5xx responses; two publish cases met latency and four cases exceeded the one-second target. All applications, jobs and profile data were removed.")

    defects = workbook["Defects"]
    defects.cell(2, 1, "DEF-001")
    defects.cell(2, 2, "ACC-02-ZM")
    defects.cell(2, 3, "Accessibility")
    defects.cell(2, 4, "Candidate profile helper text has insufficient color contrast")
    defects.cell(2, 5, "High")
    defects.cell(2, 6, "Open the candidate profile at any audited viewport and run the WCAG 2.1 A/AA axe scan.")
    defects.cell(2, 7, "Normal-size helper text meets at least the WCAG AA 4.5:1 contrast ratio.")
    defects.cell(2, 8, "Four candidate-profile text elements measured between 4.16:1 and 4.45:1; axe reported color-contrast as serious.")
    defects.cell(2, 9, "Muted text colors are too light against the profile backgrounds.")
    defects.cell(2, 10, "Pending: darken the affected muted text tokens and rerun the deployed audit.")
    defects.cell(2, 11, "Development Team")
    defects.cell(2, 12, "Open")
    defects.cell(2, 13, date(2026, 9, 10))
    defects.cell(2, 16, "EV-ACC-001")

    defects.cell(3, 1, "DEF-002")
    defects.cell(3, 2, "ACC-02-ZM")
    defects.cell(3, 3, "Accessibility")
    defects.cell(3, 4, "Candidate profile header overflows in the automated 200% zoom scenario")
    defects.cell(3, 5, "High")
    defects.cell(3, 6, "Open /candidate or /candidate/setup at the 640 px audit viewport and apply 200% CSS zoom.")
    defects.cell(3, 7, "Essential profile content reflows without horizontal scrolling.")
    defects.cell(3, 8, "Document width increased from 640 px to 832 px because the header tools/language control did not reflow.")
    defects.cell(3, 9, "The candidate profile header tools retain a non-wrapping width at high zoom.")
    defects.cell(3, 10, "Pending manual browser-zoom confirmation, then responsive header correction and retest.")
    defects.cell(3, 11, "Development Team")
    defects.cell(3, 12, "Open")
    defects.cell(3, 13, date(2026, 9, 10))
    defects.cell(3, 16, "EV-ACC-001")

    defects.cell(4, 1, "DEF-003")
    defects.cell(4, 2, "SEC-12")
    defects.cell(4, 3, "Security and Privacy")
    defects.cell(4, 4, "Oversized multipart request creates an application after PHP drops the file")
    defects.cell(4, 5, "High")
    defects.cell(4, 6, "Submit a 10 MB + 1 byte application document to the deployed candidate application endpoint.")
    defects.cell(4, 7, "The request is rejected before persistence because it exceeds the documented 10 MB limit.")
    defects.cell(4, 8, "PHP post_max_size is 8 MB, emits a startup warning and drops the multipart fields; the controller then creates an application without the rejected document and returns a success body.")
    defects.cell(4, 9, "The runtime post_max_size is lower than the application limit, and the controller does not reject an empty multipart body caused by request-size truncation.")
    defects.cell(4, 10, "Pending: align PHP request limits above 10 MB and reject truncated/oversized requests before creating an application.")
    defects.cell(4, 11, "Development Team")
    defects.cell(4, 12, "Open")
    defects.cell(4, 13, date(2026, 9, 10))
    defects.cell(4, 16, "EV-SEC-002")

    defects.cell(5, 1, "DEF-004")
    defects.cell(5, 2, "AI-EN-C08, AI-EN-I01, AI-EN-I08, AI-FR-C01, AI-FR-C08, AI-FR-I01, AI-FR-I08")
    defects.cell(5, 3, "AI Profile Suggestions")
    defects.cell(5, 4, "AI intermittently omits the explicit work-and-training opportunity preference")
    defects.cell(5, 5, "High")
    defects.cell(5, 6, "Submit clean and imperfect English or French narratives that explicitly request both paid work and hospitality training.")
    defects.cell(5, 7, "The structured opportunityPreference is always returned as both when the candidate explicitly requests both.")
    defects.cell(5, 8, "Seven of sixty gold-standard cases omitted the both preference; other expected structured fields remained available.")
    defects.cell(5, 9, "The generative extraction path does not deterministically normalize every explicit both-work-and-training phrase.")
    defects.cell(5, 10, "Pending: add deterministic post-processing for explicit both phrases in all supported languages and rerun the 60-case matrix.")
    defects.cell(5, 11, "Development Team")
    defects.cell(5, 12, "Open")
    defects.cell(5, 13, date(2026, 9, 10))
    defects.cell(5, 16, "EV-AI-002")

    defects.cell(6, 1, "DEF-005")
    defects.cell(6, 2, "PERF-08-L10, PERF-09-L1, PERF-09-L5, PERF-09-L10")
    defects.cell(6, 3, "Performance and Reliability")
    defects.cell(6, 4, "Write-operation p95 exceeds the one-second staging target")
    defects.cell(6, 5, "Medium")
    defects.cell(6, 6, "Publish or submit unique valid records for three rounds at concurrency 1, 5 and 10.")
    defects.cell(6, 7, "Each valid non-AI operation completes with p95 below one second and fewer than one percent 5xx responses.")
    defects.cell(6, 8, "All 96 writes succeeded with zero 5xx responses, but publish L10 measured 1476.6 ms and application L1/L5/L10 measured 1216.6/4825.7/10685.9 ms p95.")
    defects.cell(6, 9, "The synchronous application path includes scoring and persistence work; concurrent writes queue on the small free staging service/database.")
    defects.cell(6, 10, "Pending: profile the application transaction and scoring call, optimize database writes, then rerun on the intended staging capacity.")
    defects.cell(6, 11, "Development Team")
    defects.cell(6, 12, "Open")
    defects.cell(6, 13, date(2026, 9, 10))
    defects.cell(6, 16, "EV-PERF-005")

    workbook.save(path)


if __name__ == "__main__":
    update_workbook(WORKBOOK)
    print(f"Recorded deployed multilingual, matching, AI, workflow, performance, security and accessibility evidence in {WORKBOOK.resolve()}")
